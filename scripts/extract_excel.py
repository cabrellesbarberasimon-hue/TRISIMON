"""Extrae alimentos y la semana de ejemplo de APP_SIMON_DIETA.xlsx y genera
src/db/seed/excelData.ts. Uso: python3 scripts/extract_excel.py (requiere openpyxl).

Las macros de cada comida se reconstruyen parseando las fórmulas del Excel
(p. ej. =ROUND((200*$C$70+10*$C$74)/100,0) -> 200 g del alimento de la fila 70).
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
wb = openpyxl.load_workbook(ROOT / 'APP_SIMON_DIETA.xlsx')
wb_values = openpyxl.load_workbook(ROOT / 'APP_SIMON_DIETA.xlsx', data_only=True)
ws = wb.active
wv = wb_values.active

# fila del Excel -> (id, nombre, marca, unidad, excluido)
FOODS = {
    53: ('patata-hervida', 'Patata hervida', '', 'g', False),
    54: ('batata-hervida', 'Batata hervida', '', 'g', False),
    55: ('boniato-asado', 'Boniato asado', '', 'g', False),
    56: ('patata-cocida', 'Patata cocida', '', 'g', False),
    58: ('claras-huevo', 'Claras de huevo pasteurizadas', 'Mercadona', 'ml', False),
    59: ('lomo-embuchado', 'Lomo embuchado', 'Costa Brava', 'g', False),
    60: ('macarrones-integrales', 'Macarrones integrales (crudo)', 'Hacendado', 'g', False),
    61: ('arroz-redondo', 'Arroz redondo (crudo)', 'Hacendado', 'g', False),
    62: ('arroz-microondas', 'Arroz microondas (cocido, bolsita 167 g)', 'Hacendado', 'g', False),
    63: ('tomate-triturado', 'Tomate triturado', 'Hacendado', 'g', False),
    64: ('atun-natural', 'Atún claro al natural (escurrido, lata 60 g)', '', 'g', False),
    65: ('pechuga-pollo', 'Pechuga de pollo', 'Genérico', 'g', False),
    66: ('aove', 'AOVE (aceite de oliva virgen extra)', 'Genérico', 'g', False),
    67: ('pechuga-pavo-fresca', 'Pechuga de pavo fresca', 'Genérico', 'g', False),
    68: ('pan-integral', 'Pan integral 100% (pieza 95 g)', 'Consum', 'g', False),
    69: ('queso-madurado', 'Queso mezcla madurado', 'Entrepinares', 'g', False),
    70: ('leche-semidesnatada', 'Leche semidesnatada', 'Central Lechera Asturiana', 'ml', False),
    71: ('pavo-92', 'Pechuga de pavo 92%', 'Hacendado', 'g', False),
    72: ('platano', 'Plátano (sin piel)', 'Genérico', 'g', False),
    73: ('nueces', 'Nueces', '', 'g', True),
    74: ('miel', 'Miel', 'Genérico', 'g', False),
    75: ('copos-avena', 'Copos de avena', '', 'g', True),
    76: ('naranja', 'Naranja (sin piel)', 'Genérico', 'g', False),
    77: ('tomate-natural', 'Tomate natural', 'Genérico', 'g', False),
}
SOURCE_BY_NAME = {ws.cell(r, 2).value: FOODS[r][0] for r in (53, 54, 55, 56)}

foods = []
for row, (fid, name, brand, unit, excluded) in FOODS.items():
    kcal, p, c, g = (ws.cell(row, col).value for col in (3, 4, 5, 6))
    foods.append({
        'id': fid, 'name': name, 'brand': brand, 'unit': unit,
        'kcal': kcal, 'p': p, 'c': c, 'g': g, 'excluded': excluded,
        'notes': 'NO TOMAR' if excluded else '',
    })

DAY_COLS = [3, 8, 13, 18, 23, 28, 33]  # C, H, M, R, W, AB, AG
SLOTS = {17: 'desayuno', 18: 'mediaManana', 19: 'comida', 20: 'merienda'}
DAY_TYPE_IDS = {
    'DESCANSO / MUY SUAVE': 'descanso',
    'ENTRENAMIENTO NORMAL': 'normal',
    'DOBLE SESIÓN': 'doble',
    'GRAN CARGA / BICI LARGA': 'gran-carga',
}
TERM = re.compile(r'(\d+(?:\.\d+)?)\*\$C\$(\d+)')


def split_label(text: str) -> str:
    first = text.split('\n')[0].strip()
    return '' if first.startswith('·') else first


def clean(text: str) -> str:
    parts = [ln.strip().lstrip('·').strip() for ln in text.split('\n')]
    return ' + '.join(p for p in parts if p)


days = []
for i, col in enumerate(DAY_COLS):
    meals = {}
    for row, slot in SLOTS.items():
        text = ws.cell(row, col).value
        formula = ws.cell(row, col + 1).value
        lines = [{'kind': 'food', 'foodId': FOODS[int(r)][0], 'grams': float(gr)}
                 for gr, r in TERM.findall(formula)]
        meals[slot] = {'label': split_label(text), 'lines': lines}

    plate = ws.cell(44, col).value
    fixed = ws.cell(45, col).value
    fixed_line = {
        'kind': 'free',
        'name': 'Parte fija (Excel): ' + clean('\n'.join(plate.split('\n')[1:] if split_label(plate) else plate.split('\n')) + '\n' + fixed),
        'kcal': ws.cell(45, col + 1).value, 'p': ws.cell(45, col + 2).value,
        'c': ws.cell(45, col + 3).value, 'g': ws.cell(45, col + 4).value,
    }
    source_id = SOURCE_BY_NAME[ws.cell(43, col).value]
    grams = wv.cell(49, col + 3).value
    meals['cena'] = {
        'label': split_label(plate) or 'CENA',
        'lines': [fixed_line, {'kind': 'food', 'foodId': source_id, 'grams': grams}],
    }
    meals['entreno'] = {'label': '', 'lines': []}
    days.append({
        'dayTypeId': DAY_TYPE_IDS[ws.cell(15, col).value],
        'meals': meals,
        'dinner': {'sourceFoodId': source_id, 'fixedLines': [fixed_line]},
        # Valores de control calculados por Excel (para los tests de la fase 2)
        'excelCheck': {
            'totals': [wv.cell(22, col + k).value for k in range(1, 5)],
            'semaphore': [wv.cell(25, col + k).value for k in range(1, 5)],
            'dinnerGrams': grams,
        },
    })

day_types = []
for row in range(8, 12):
    name = ws.cell(row, 2).value
    v = [ws.cell(row, c).value for c in range(3, 11)]
    day_types.append({
        'id': DAY_TYPE_IDS[name], 'name': name, 'order': row - 8,
        'kcal': [v[0], v[1]], 'p': [v[2], v[3]], 'c': [v[4], v[5]], 'g': [v[6], v[7]],
    })

out = {
    'methodology': ws['C33'].value,
    'foods': foods,
    'dayTypes': day_types,
    'weekStart': '2026-08-31',
    'days': days,
}
target = ROOT / 'src/db/seed/excelData.ts'
target.write_text(
    '// Generado por scripts/extract_excel.py desde APP_SIMON_DIETA.xlsx. No editar a mano.\n'
    'import type { ExcelSeed } from \'./types\';\n\n'
    'export const excelData: ExcelSeed = ' + json.dumps(out, ensure_ascii=False, indent=2) + ';\n',
    encoding='utf-8',
)
print('OK', target)
